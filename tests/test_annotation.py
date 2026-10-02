"""Выборка для ручной разметки, файлы Excel и метрики качества. Сеть не нужна."""
import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from cryptonews import annotation


def make_news(n=3000, seed=0, sources=("a.com", "b.com", "c.com")) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    times = pd.to_datetime(rng.integers(pd.Timestamp("2018-01-01").value // 10**9,
                                        pd.Timestamp("2025-07-31").value // 10**9, n), unit="s", utc=True)
    weights = np.array([0.6, 0.3, 0.1])[:len(sources)]
    return pd.DataFrame({
        "published_utc": times,
        "source": rng.choice(list(sources), size=n, p=weights / weights.sum()),
        "title": [f"Headline number {i}" for i in range(n)],
        "url": [f"https://site/{i}" for i in range(n)],
    })


def test_allocate_is_proportional_and_exact():
    sizes = pd.Series({"x": 500, "y": 300, "z": 3, "w": 197})
    quotas = annotation.allocate(sizes, 100)
    assert quotas.sum() == 100
    assert (quotas <= sizes).all()
    assert quotas["x"] == 50 and quotas["y"] == 30
    assert annotation.allocate(pd.Series({"x": 2, "y": 1}), 10).sum() == 3


def test_stratified_sample_is_deterministic_and_proportional():
    news = make_news()
    news["year"] = news["published_utc"].dt.year
    first = annotation.stratified_sample(news, 400, ["year", "source"], seed=7)
    again = annotation.stratified_sample(news.sample(frac=1, random_state=1), 400, ["year", "source"], seed=7)
    assert len(first) == 400 and first["url"].is_unique
    assert first["url"].tolist() == again["url"].tolist()      # порядок строк на входе не важен
    share = first["source"].value_counts(normalize=True)
    assert abs(share["a.com"] - 0.6) < 0.05


def test_make_samples_ids_and_second_subset():
    news = make_news()
    main, second = annotation.make_samples(news, "en", 400, 100, seed=2026)
    assert main["id"].tolist()[:2] == ["EN-001", "EN-002"] and main["id"].is_unique
    assert set(second["id"]) <= set(main["id"]) and len(second) == 100
    assert second["id"].is_monotonic_increasing
    assert list(main.columns) == annotation.KEY_COLUMNS
    assert main["published_utc"].str.endswith("+0000").all()
    ru_main, _ = annotation.make_samples(news.assign(source="forklog.com"), "ru", 50, 10, seed=1)
    assert ru_main["stratum"].str.match(r"^\d{4}$").all()      # у одного источника — только годы
    table = annotation.strata_table(main, news)
    assert table["В выборке"].sum() == 400


def test_workbook_roundtrip(tmp_path):
    news = make_news(500)
    news.loc[0, "title"] = "Bitcoin \x0bsurges"                # недопустимый в xlsx символ
    main, _ = annotation.make_samples(news, "en", 20, 5, seed=3)
    path = tmp_path / "annotation_en_main.xlsx"
    annotation.write_workbook(main, path, "# Инструкция\n\nТекст.\n\n| Заголовок | Тональность |\n| --- | --- |\n| A | позитивная |\n", "Заголовок файла")

    book = load_workbook(path)
    assert book.sheetnames == ["Инструкция", "Разметка", "_key"]
    assert book["_key"].sheet_state == "hidden"
    sheet = book["Разметка"]
    assert sheet["A2"].value == main.loc[0, "id"] and sheet["C2"].value is None
    assert sheet.data_validations.dataValidation[0].sqref is not None
    marks = ["позитивная", "нейтральная", "негативная", "+", "0", "-"]
    for row, mark in enumerate(marks, start=2):
        sheet.cell(row=row, column=3, value=mark)
    sheet.cell(row=3, column=4, value="непонятно")
    book.save(path)

    labels = annotation.read_workbook(path)
    assert labels["label"].head(6).tolist() == ["positive", "neutral", "negative", "positive", "neutral", "negative"]
    assert labels["label"].isna().sum() == 14
    assert labels.loc[1, "comment"] == "непонятно"
    assert not labels["title_changed"].any()
    assert labels["url"].tolist() == main["url"].tolist()


def test_parse_label_rejects_garbage():
    assert annotation.parse_label(" Позитивная ") == "positive"
    assert annotation.parse_label(None) is None and annotation.parse_label("  ") is None
    with pytest.raises(ValueError, match="Непонятная"):
        annotation.parse_label("хорошая")


def test_metrics_match_sklearn_and_bootstrap_is_paired():
    from sklearn.metrics import f1_score

    rng = np.random.default_rng(0)
    classes = np.array(["negative", "neutral", "positive"])
    truth = rng.choice(classes, size=300, p=[0.25, 0.5, 0.25])
    good = np.where(rng.random(300) < 0.8, truth, rng.choice(classes, size=300))
    bad = rng.choice(classes, size=300)
    metrics = annotation.classification_metrics(truth, good)
    assert abs(metrics["macro_f1"] - f1_score(truth, good, average="macro")) < 1e-12
    codes_t, codes_p = annotation._codes(truth), annotation._codes(good)
    counts = np.zeros((3, 3), dtype=int)
    np.add.at(counts, (codes_t, codes_p), 1)
    assert abs(annotation._macro_f1_from_counts(counts) - metrics["macro_f1"]) < 1e-12
    boot = annotation.bootstrap_macro_f1(truth, {"good": good, "bad": bad}, reps=500, seed=1)
    assert boot.shape == (500, 2)
    assert np.percentile(boot[:, 0] - boot[:, 1], 2.5) > 0
    assert abs(boot[:, 0].mean() - metrics["macro_f1"]) < 0.02


def test_agreement_kappa():
    a = ["positive", "neutral", "negative", "neutral"] * 25
    result = annotation.agreement(a, a)
    assert abs(result["cohen_kappa"] - 1.0) < 1e-12 and result["percent_agreement"] == 1.0
    b = a[1:] + a[:1]
    assert annotation.agreement(a, b)["cohen_kappa"] < 0.5
