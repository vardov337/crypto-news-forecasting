"""Ручная разметка тональности (задачи 5.1–5.4, PROTOCOL.md, раздел 5).

Выборка заголовков стратифицирована: по годам и источникам для англоязычных
новостей, по годам — для ForkLog. Объём страты пропорционален числу новостей в ней,
поэтому качество модели на разметке относится ко всему корпусу, который идёт
в признаки. Новости берутся только из периода выборки.

Разметчик видит только номер и заголовок: ни источника, ни даты, ни оценок моделей.
Адрес и служебные поля лежат на скрытом листе: по номеру разметка связывается
с новостью, даже если строки на основном листе переставят.
"""
from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd

from cryptonews.config import PROJECT_ROOT
from cryptonews.sentiment import CLASSES

LABELS_RU = {"позитивная": "positive", "нейтральная": "neutral", "негативная": "negative"}
ALIASES = {
    **LABELS_RU,
    "позитив": "positive", "позитивно": "positive", "поз": "positive", "+": "positive",
    "нейтрально": "neutral", "нейтрал": "neutral", "нейтр": "neutral", "0": "neutral",
    "негатив": "negative", "негативно": "negative", "нег": "negative",
    "-": "negative", "−": "negative", "–": "negative",
    "positive": "positive", "neutral": "neutral", "negative": "negative",
}
KEY_COLUMNS = ["id", "url", "title", "published_utc", "source", "year", "stratum"]
LABEL_COLUMNS = ["id", "url", "title", "label", "comment"]
SHEET_HELP, SHEET_MAIN, SHEET_GLOSSARY, SHEET_KEY = "Инструкция", "Разметка", "Словарь", "_key"
HEADER_ID, HEADER_TITLE, HEADER_LABEL, HEADER_COMMENT = "№", "Заголовок", "Тональность", "Комментарий (необязательно)"
HEADER_TRANSLATION = "Перевод на русский (подсказка)"
FILLS = {"позитивная": "C6EFCE", "нейтральная": "E7E6E6", "негативная": "FFC7CE"}


# --- Выборка

def allocate(sizes: pd.Series, total: int) -> pd.Series:
    """Пропорциональное распределение total по стратам методом наибольших остатков.

    Квота страты не превышает её размера, сумма квот равна total (или числу
    всех записей, если их меньше)."""
    sizes = sizes.astype(int)
    total = min(int(total), int(sizes.sum()))
    quota = sizes / sizes.sum() * total
    base = np.floor(quota).astype(int).clip(upper=sizes)
    remainder = total - int(base.sum())
    for key in (quota - base).sort_values(ascending=False, kind="stable").index:
        if remainder <= 0:
            break
        if base[key] < sizes[key]:
            base[key] += 1
            remainder -= 1
    return base


def stratified_sample(frame: pd.DataFrame, total: int, strata: list[str], seed: int) -> pd.DataFrame:
    """Случайная выборка без возвращения с пропорциональным размещением по стратам."""
    data = frame.sort_values(["published_utc", "url"]).reset_index(drop=True)
    data["stratum"] = data[strata].astype(str).agg(" | ".join, axis=1)
    groups = data.groupby("stratum", sort=True)
    sizes = groups.size()
    quotas = allocate(sizes, total)
    positions = groups.indices
    rng = np.random.default_rng(seed)
    picked: list[int] = []
    for key in sizes.index:
        n = int(quotas[key])
        if n:
            picked.extend(rng.choice(positions[key], size=n, replace=False).tolist())
    return data.iloc[sorted(picked)].reset_index(drop=True)


def make_samples(news: pd.DataFrame, language: str, size: int, second_size: int,
                 seed: int) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Основная выборка для разметки и её часть для второго разметчика.

    Строки перемешаны, чтобы порядок не выдавал год и источник; номера вида EN-001."""
    data = news.assign(year=news["published_utc"].dt.year)
    strata = ["year", "source"] if data["source"].nunique() > 1 else ["year"]
    sample = stratified_sample(data, size, strata, seed)
    rng = np.random.default_rng(seed + 1)
    sample = sample.iloc[rng.permutation(len(sample))].reset_index(drop=True)
    width = max(3, len(str(len(sample))))
    sample.insert(0, "id", [f"{language.upper()}-{i + 1:0{width}d}" for i in range(len(sample))])
    sample["published_utc"] = sample["published_utc"].dt.strftime("%Y-%m-%d %H:%M:%S%z")
    second_rows = np.sort(rng.choice(len(sample), size=min(second_size, len(sample)), replace=False))
    second = sample.iloc[second_rows].reset_index(drop=True)
    return sample[KEY_COLUMNS], second[KEY_COLUMNS]


def strata_table(sample: pd.DataFrame, news: pd.DataFrame) -> pd.DataFrame:
    """Сравнение структуры выборки и корпуса — для приложения к статье."""
    data = news.assign(year=news["published_utc"].dt.year)
    strata = ["year", "source"] if data["source"].nunique() > 1 else ["year"]
    corpus = data.groupby(strata).size().rename("В корпусе")
    taken = sample.groupby(strata).size().rename("В выборке")
    table = pd.concat([corpus, taken], axis=1).fillna(0).astype(int)
    table["Доля корпуса"] = (table["В корпусе"] / table["В корпусе"].sum()).round(4)
    table["Доля выборки"] = (table["В выборке"] / max(table["В выборке"].sum(), 1)).round(4)
    return table


# --- Файлы для разметчиков

def instruction_rows(text: str) -> list[tuple[list[str], bool]]:
    """Строки инструкции для листа Excel: разметка Markdown убирается, таблица — в две колонки."""
    rows: list[tuple[list[str], bool]] = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped or re.fullmatch(r"\|?\s*:?-{3,}.*", stripped):
            continue                                   # пустые строки и разделитель таблицы
        if stripped.startswith("|"):
            cells = [c.strip().replace("**", "") for c in stripped.strip("|").split("|")]
            rows.append((cells, cells[0] == "Заголовок"))
            continue
        heading = stripped.startswith("#")
        clean = re.sub(r"^#+\s*", "", stripped).replace("**", "").replace("`", "")
        clean = re.sub(r"^-\s+", "• ", clean)
        if heading and rows:
            rows.append(([""], False))                 # отступ перед заголовком раздела
        rows.append(([clean], heading))
    return rows


def excel_safe(text) -> str:
    """Текст без управляющих символов, которые формат xlsx не допускает."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    return ILLEGAL_CHARACTERS_RE.sub("", str(text))


def load_translations(language: str) -> pd.DataFrame | None:
    """Вспомогательный перевод заголовков для разметчика (annotation/translations_<язык>.csv)."""
    path = PROJECT_ROOT / "annotation" / f"translations_{language}.csv"
    return pd.read_csv(path, dtype=str) if path.exists() else None


def load_glossary(language: str) -> pd.DataFrame | None:
    """Словарь терминов для разметчика (annotation/glossary_<язык>.csv): term, meaning."""
    path = PROJECT_ROOT / "annotation" / f"glossary_{language}.csv"
    if not path.exists():
        return None
    glossary = pd.read_csv(path, dtype=str)
    return glossary.sort_values("term", key=lambda t: t.str.lower()).reset_index(drop=True)


def add_translations(sample: pd.DataFrame, translations: pd.DataFrame | None) -> pd.DataFrame:
    """Добавляет к выборке колонку translation (по адресу новости), порядок строк сохраняется."""
    if translations is None:
        return sample
    mapping = translations.drop_duplicates("url").set_index("url")["translation"]
    return sample.assign(translation=sample["url"].map(mapping))


def write_workbook(sample: pd.DataFrame, path: Path, instruction: str, heading: str,
                   description: str | None = None, glossary: pd.DataFrame | None = None) -> None:
    """Файл Excel для разметчика: инструкция, лист разметки, словарь терминов и скрытый лист ключей.

    Если в выборке есть колонка translation, рядом с заголовком появляется перевод-подсказка."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.properties import PageSetupProperties

    sample = sample.assign(title=sample["title"].map(excel_safe))
    translated = "translation" in sample.columns and sample["translation"].notna().any()
    book = Workbook()
    help_sheet = book.active
    help_sheet.title = SHEET_HELP
    help_sheet.column_dimensions["A"].width = 95
    help_sheet.column_dimensions["B"].width = 18
    help_sheet.append([heading])
    help_sheet["A1"].font = Font(bold=True, size=13)
    for cells, bold in instruction_rows(instruction):
        help_sheet.append(cells)
        for cell in help_sheet[help_sheet.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if bold:
                cell.font = Font(bold=True)

    sheet = book.create_sheet(SHEET_MAIN)
    headers = [HEADER_ID, HEADER_TITLE] + ([HEADER_TRANSLATION] if translated else []) + [HEADER_LABEL, HEADER_COMMENT]
    sheet.append(headers)
    for row in sample.itertuples(index=False):
        hint = ([excel_safe(row.translation) if pd.notna(row.translation) else None] if translated else [])
        sheet.append([row.id, row.title, *hint, None, None])
    last = len(sample) + 1
    letter = {name: get_column_letter(i + 1) for i, name in enumerate(headers)}
    label, comment = letter[HEADER_LABEL], letter[HEADER_COMMENT]
    counter = get_column_letter(len(headers) + 2)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    widths = {letter[HEADER_ID]: 10, letter[HEADER_TITLE]: 60 if translated else 85,
              label: 16, comment: 30, counter: 24}
    if translated:
        widths[letter[HEADER_TRANSLATION]] = 60
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    wrapped = {letter[HEADER_TITLE], comment} | ({letter[HEADER_TRANSLATION]} if translated else set())
    for row in sheet.iter_rows(min_row=2, max_row=last, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.alignment = Alignment(wrap_text=cell.column_letter in wrapped, vertical="top")
            if translated and cell.column_letter == letter[HEADER_TRANSLATION]:
                cell.font = Font(color="404040", italic=True)
    sheet.freeze_panes = "A2"
    choices = DataValidation(
        type="list", formula1='"' + ",".join(LABELS_RU) + '"', allow_blank=True,
        showErrorMessage=True, errorTitle="Неизвестное значение",
        error="Выберите из списка: позитивная, нейтральная или негативная",
    )
    sheet.add_data_validation(choices)
    choices.add(f"{label}2:{label}{last}")
    for word, color in FILLS.items():
        sheet.conditional_formatting.add(
            f"{label}2:{label}{last}",
            CellIsRule(operator="equal", formula=[f'"{word}"'],
                       fill=PatternFill(start_color=color, end_color=color, fill_type="solid")))
    sheet[f"{counter}1"] = "Размечено"
    sheet[f"{counter}1"].font = Font(bold=True)
    sheet[f"{counter}2"] = f'=COUNTA({label}2:{label}{last})&" из {len(sample)}"'
    pages = [(help_sheet, False), (sheet, True)]

    if glossary is not None and len(glossary):
        terms = book.create_sheet(SHEET_GLOSSARY)
        terms.append(["Термин", "Значение"])
        for cell in terms[1]:
            cell.font = Font(bold=True)
        for item in glossary.itertuples(index=False):
            terms.append([item.term, item.meaning])
        terms.column_dimensions["A"].width = 34
        terms.column_dimensions["B"].width = 100
        for row in terms.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        terms.freeze_panes = "A2"
        pages.append((terms, False))

    for page, landscape in pages:
        page.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        page.page_setup.fitToWidth, page.page_setup.fitToHeight = 1, 0
        page.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.print_title_rows = "1:1"
    sheet.print_area = f"A1:{comment}{last}"

    key = book.create_sheet(SHEET_KEY)
    key.append(KEY_COLUMNS)
    for row in sample[KEY_COLUMNS].itertuples(index=False):
        key.append([None if pd.isna(v) else (int(v) if isinstance(v, np.integer) else v) for v in row])
    key.sheet_state = "hidden"
    if description:
        book.properties.description = description      # период и зерно — чтобы опознать версию файла
    book.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def parse_label(value) -> str | None:
    """Метка разметчика → negative / neutral / positive; пустая ячейка → None."""
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return None
    key = str(value).strip().lower()
    if not key:
        return None
    if key not in ALIASES:
        raise ValueError(f"Непонятная метка «{value}»: ожидается позитивная, нейтральная или негативная")
    return ALIASES[key]


def read_workbook(path: Path) -> pd.DataFrame:
    """Размеченный файл → таблица с адресом, заголовком, меткой и комментарием.

    Колонки ищутся по названиям (в файлах с переводом их порядок другой); строки
    связываются по номеру через скрытый лист; изменённые заголовки помечаются."""
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True)
    rows = list(book[SHEET_MAIN].iter_rows(values_only=True))
    header = ["" if v is None else str(v).strip() for v in rows[0]]
    missing = [name for name in (HEADER_ID, HEADER_TITLE, HEADER_LABEL) if name not in header]
    if missing:
        raise ValueError(f"В файле {Path(path).name} на листе «{SHEET_MAIN}» нет колонок: {', '.join(missing)}")
    position = {name: header.index(name) for name in (HEADER_ID, HEADER_TITLE, HEADER_LABEL, HEADER_COMMENT)
                if name in header}

    def cell(row, name):
        i = position.get(name)
        return row[i] if i is not None and i < len(row) else None

    marks = pd.DataFrame(
        [{"id": cell(r, HEADER_ID), "title_in_sheet": cell(r, HEADER_TITLE),
          "label_raw": cell(r, HEADER_LABEL), "comment": cell(r, HEADER_COMMENT)}
         for r in rows[1:] if r and cell(r, HEADER_ID) is not None],
        columns=["id", "title_in_sheet", "label_raw", "comment"],
    )
    key_rows = list(book[SHEET_KEY].iter_rows(values_only=True))
    key = pd.DataFrame(key_rows[1:], columns=key_rows[0])
    merged = key.merge(marks, on="id", how="left", validate="one_to_one")
    merged["label"] = merged["label_raw"].map(parse_label)
    merged["title_changed"] = (merged["title_in_sheet"].astype("string").str.strip()
                               != merged["title"].astype("string").str.strip())
    merged["comment"] = merged["comment"].astype("string")
    return merged


# --- Качество моделей и согласие разметчиков

CLASS_INDEX = {c: i for i, c in enumerate(CLASSES)}


def _codes(labels) -> np.ndarray:
    return np.asarray([CLASS_INDEX[v] for v in labels], dtype=np.int64)


def _macro_f1_from_counts(counts: np.ndarray) -> np.ndarray:
    """macro-F1 по матрицам ошибок (… × 3 × 3, строки — истина, столбцы — прогноз)."""
    tp = np.diagonal(counts, axis1=-2, axis2=-1).astype(float)
    fp = counts.sum(axis=-2) - tp
    fn = counts.sum(axis=-1) - tp
    denom = 2 * tp + fp + fn
    f1 = np.divide(2 * tp, denom, out=np.zeros_like(tp), where=denom > 0)
    return f1.mean(axis=-1)


def classification_metrics(y_true, y_pred) -> dict:
    """accuracy, macro-F1 и F1 по классам."""
    from sklearn.metrics import accuracy_score, f1_score

    labels = list(CLASSES)
    per_class = f1_score(y_true, y_pred, labels=labels, average=None, zero_division=0)
    return {"n": len(y_true), "accuracy": float(accuracy_score(y_true, y_pred)),
            "macro_f1": float(f1_score(y_true, y_pred, labels=labels, average="macro", zero_division=0)),
            **{f"f1_{c}": float(v) for c, v in zip(labels, per_class)}}


def bootstrap_macro_f1(y_true, predictions: dict[str, list], reps: int = 2000,
                       seed: int = 0) -> np.ndarray:
    """Бутстреп-распределение macro-F1: массив reps × число моделей.

    Для всех моделей используются одни и те же выборки заголовков, поэтому разности
    между столбцами — парный бутстреп."""
    truth = _codes(y_true)
    n = len(truth)
    rng = np.random.default_rng(seed)
    idx = rng.integers(0, n, size=(reps, n))
    offsets = (np.arange(reps) * 9)[:, None]
    result = np.empty((reps, len(predictions)))
    for j, labels in enumerate(predictions.values()):
        codes = truth[idx] * 3 + _codes(labels)[idx] + offsets
        counts = np.bincount(codes.ravel(), minlength=9 * reps).reshape(reps, 3, 3)
        result[:, j] = _macro_f1_from_counts(counts)
    return result


def agreement(first, second) -> dict:
    """Согласие двух разметчиков: доля совпадений и каппа Коэна."""
    from sklearn.metrics import cohen_kappa_score, confusion_matrix

    first, second = list(first), list(second)
    matrix = confusion_matrix(first, second, labels=list(CLASSES))
    return {"n": len(first), "percent_agreement": float(np.mean(np.asarray(first) == np.asarray(second))),
            "cohen_kappa": float(cohen_kappa_score(first, second, labels=list(CLASSES))),
            "confusion": matrix.tolist()}
