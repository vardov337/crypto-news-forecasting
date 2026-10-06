"""Оформление таблиц и рисунков статьи (шаг 9).

Цвета — проверенная на различимость при нарушениях цветового зрения палитра (первые три
категориальных цвета); подписи — тёмным, сетка — едва заметная. Числа на осях — с
десятичной запятой, как принято в русскоязычных журналах.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd

INK, INK2, MUTED, GRID, AXIS = "#0b0b0b", "#52514e", "#898781", "#e1e0d9", "#c3c2b7"
SERIES = ("#2a78d6", "#eb6834", "#1baf7a")          # синий, оранжевый, бирюзовый
MONTHS = ("янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек")


def pyplot():
    """matplotlib без экрана, со стилем рисунков статьи."""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 8.5, "axes.titlesize": 9.5, "axes.labelsize": 8.5,
        "axes.titlecolor": INK, "axes.labelcolor": INK2, "axes.edgecolor": AXIS, "axes.linewidth": 0.8,
        "xtick.color": AXIS, "ytick.color": AXIS, "xtick.labelcolor": INK2, "ytick.labelcolor": INK2,
        "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.6, "grid.linestyle": "-",
        "axes.spines.top": False, "axes.spines.right": False, "axes.axisbelow": True,
        "legend.frameon": False, "legend.fontsize": 8, "lines.linewidth": 1.5,
        "lines.solid_capstyle": "round", "lines.solid_joinstyle": "round",
        "figure.facecolor": "white", "axes.facecolor": "white", "savefig.facecolor": "white",
    })
    return plt


def comma_number(value: float, decimals: int) -> str:
    text = f"{value:,.{decimals}f}".replace(",", " ").replace(".", ",")
    return text.replace("-", "−")


def comma_formatter(decimals: int = 1):
    from matplotlib.ticker import FuncFormatter

    return FuncFormatter(lambda value, _: comma_number(value, decimals))


def month_formatter():
    from matplotlib.dates import num2date
    from matplotlib.ticker import FuncFormatter

    return FuncFormatter(lambda value, _: f"{MONTHS[num2date(value).month - 1]} {num2date(value).year}")


def save_figure(fig, path_without_suffix: Path) -> list[Path]:
    paths = []
    for suffix, kwargs in ((".png", {"dpi": 300}), (".svg", {})):
        path = path_without_suffix.with_suffix(suffix)
        fig.savefig(path, bbox_inches="tight", **kwargs)
        paths.append(path)
    return paths


def write_sheets(path: Path, sheets: list[tuple[str, list[tuple[str, pd.DataFrame]]]]) -> None:
    """Книга Excel: на каждом листе одна или несколько таблиц, у каждой — заголовок строкой выше."""
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, blocks in sheets:
            row = 0
            widths: dict[int, int] = {}
            for title, frame in blocks:
                frame.to_excel(writer, sheet_name=name, startrow=row + 1, index=False)
                sheet = writer.sheets[name]
                sheet.cell(row=row + 1, column=1, value=title).font = Font(bold=True, color="0B0B0B")
                for j, column in enumerate(frame.columns, start=1):
                    cell = sheet.cell(row=row + 2, column=j)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    values = [str(column)] + [str(v) for v in frame[column].tolist()]
                    widths[j] = max(widths.get(j, 8), min(48, max(len(v) for v in values) + 2))
                row += len(frame) + 3
            for j, width in widths.items():
                writer.sheets[name].column_dimensions[get_column_letter(j)].width = width
